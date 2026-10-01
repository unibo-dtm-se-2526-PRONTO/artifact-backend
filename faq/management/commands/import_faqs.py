"""Seed the FAQ knowledge base from the helpdesk's spreadsheet.

The sheet is the phone helpdesk's log, one question per row, with the columns
``Ufficio | Office | Domanda | Question | Risposta | Answer`` (office,
question and answer, each in Italian and English). ``Question`` and ``Answer``
are optional. Columns are found by their heading, so their order does not
matter.

What happens to each row:

- the office is looked up by its Italian name, then its English one; a row
  whose office is not one of ``OfficeCode`` is skipped, as is a row without a
  question or an answer. Skipped rows are counted by reason, never fatal
- every text is anonymised (``faq.anonymisation``) before it is stored or
  compared, so personal data never reaches the database
- a missing English answer is replaced by the Italian one until someone
  translates it in the admin; a missing English question falls back to the
  Italian one the same way
- new FAQs are created unpublished, because the anonymisation of names is a
  heuristic and the text is about to be shown on a public endpoint: someone
  reads them in the admin first. ``--publish`` skips that review

The import is idempotent. ``(office_code, question_it)`` is the natural key
(see ``Faq.Meta``), so running it again updates the FAQs it finds instead of
duplicating them: the answers and the English question follow the sheet. When
the sheet has no English answer, the stored one is replaced only while it is
still the untranslated copy. Whether a FAQ is published is never touched. The whole
run is one transaction: a failure halfway stores nothing.
"""

from collections import Counter
from pathlib import Path
from zipfile import BadZipFile

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from faq.anonymisation import anonymise
from faq.models import Faq
from pronto.enums import OfficeCode

# The office names used by the helpdesk, in both languages, lower-cased. The
# codes and labels of ``OfficeCode`` are accepted as well.
OFFICE_NAMES = {
    "orientamento": OfficeCode.GUIDANCE,
    "guidance": OfficeCode.GUIDANCE,
    "guidance and admissions": OfficeCode.GUIDANCE,
    "segreteria": OfficeCode.ADMIN_OFFICE,
    "segreteria studenti": OfficeCode.ADMIN_OFFICE,
    "student administration office": OfficeCode.ADMIN_OFFICE,
    "relazioni internazionali": OfficeCode.INTERNATIONAL,
    "international relations": OfficeCode.INTERNATIONAL,
    "tirocini": OfficeCode.INTERNSHIPS,
    "internships": OfficeCode.INTERNSHIPS,
    **{value.lower(): OfficeCode(value) for value, _ in OfficeCode.choices},
    **{str(label).lower(): OfficeCode(value) for value, label in OfficeCode.choices},
}

REQUIRED_COLUMNS = ("Domanda", "Risposta")
OFFICE_COLUMNS = ("Ufficio", "Office")

QUESTION_MAX_LENGTH = Faq._meta.get_field("question_it").max_length


class SkippedRow(Exception):
    """A row that cannot become a FAQ; the message is the reason."""


class Command(BaseCommand):
    help = (
        "Import FAQs from the helpdesk spreadsheet (.xlsx), anonymising them. "
        "Running it again updates the FAQs instead of duplicating them."
    )

    def add_arguments(self, parser):
        parser.add_argument("path", help="The .xlsx file to import.")
        parser.add_argument(
            "--sheet", help="The sheet to read; the first one by default."
        )
        parser.add_argument(
            "--publish",
            action="store_true",
            help="Publish new FAQs at once instead of leaving them for review.",
        )

    def handle(self, *args, path, sheet=None, publish=False, **options):
        rows = read_rows(Path(path), sheet)
        outcomes = Counter()
        skipped = Counter()
        with transaction.atomic():
            for row in rows:
                try:
                    outcomes[import_row(row, publish)] += 1
                except SkippedRow as reason:
                    skipped[str(reason)] += 1
        self.report(outcomes, skipped)

    def report(self, outcomes, skipped):
        self.stdout.write(f"Created: {outcomes['created']}")
        self.stdout.write(f"Updated: {outcomes['updated']}")
        self.stdout.write(f"Unchanged: {outcomes['unchanged']}")
        self.stdout.write(f"Skipped: {skipped.total()}")
        for reason, count in sorted(skipped.items()):
            self.stdout.write(f"  {reason}: {count}")


def read_rows(path, sheet_name):
    """The data rows of the sheet, as dicts keyed by lower-cased heading."""
    if not path.is_file():
        raise CommandError(f"File not found: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile) as error:
        raise CommandError(f"Not an .xlsx workbook: {path} ({error})") from None
    try:
        if sheet_name is None:
            return sheet_records(workbook.worksheets[0])
        if sheet_name not in workbook.sheetnames:
            raise CommandError(f"No sheet named {sheet_name!r} in {path}")
        return sheet_records(workbook[sheet_name])
    finally:
        workbook.close()


def sheet_records(sheet):
    rows = sheet.iter_rows(values_only=True)
    header = [cell_text(cell) for cell in next(rows, ())]
    columns = {name.casefold(): index for index, name in enumerate(header) if name}
    missing = [name for name in REQUIRED_COLUMNS if name.casefold() not in columns]
    if not any(name.casefold() in columns for name in OFFICE_COLUMNS):
        missing.append(" or ".join(OFFICE_COLUMNS))
    if missing:
        raise CommandError(f"Missing column(s): {', '.join(missing)}")
    return [
        {
            name: cell_text(row[index]) if index < len(row) else ""
            for name, index in columns.items()
        }
        for row in rows
    ]


def cell_text(value):
    return "" if value is None else str(value).strip()


def import_row(row, publish):
    """Create or update the FAQ a row describes: "created", "updated" or
    "unchanged". Raises ``SkippedRow`` when the row cannot become one."""
    if not any(row.values()):
        raise SkippedRow("empty row")
    office_code = office_of(row)
    if not row.get("domanda"):
        raise SkippedRow("no question")
    if not row.get("risposta"):
        raise SkippedRow("no answer")

    question_it = anonymise(row["domanda"])
    question_en = anonymise(row.get("question", "")) or question_it
    answer_it = anonymise(row["risposta"])
    sheet_answer_en = anonymise(row.get("answer", ""))
    if max(len(question_it), len(question_en)) > QUESTION_MAX_LENGTH:
        raise SkippedRow(f"question longer than {QUESTION_MAX_LENGTH} characters")

    faq = Faq.objects.filter(office_code=office_code, question_it=question_it).first()
    if faq is None:
        Faq.objects.create(
            office_code=office_code,
            question_it=question_it,
            question_en=question_en,
            answer_it=answer_it,
            answer_en=sheet_answer_en or answer_it,
            is_active=publish,
        )
        return "created"

    if sheet_answer_en:
        answer_en = sheet_answer_en
    elif faq.answer_en == faq.answer_it:
        answer_en = answer_it
    else:
        # An English answer that differs from the Italian one was written by
        # someone in the admin: the sheet has nothing better to replace it with.
        answer_en = faq.answer_en
    changes = {
        "question_en": question_en,
        "answer_it": answer_it,
        "answer_en": answer_en,
    }
    if all(getattr(faq, field) == value for field, value in changes.items()):
        return "unchanged"
    for field, value in changes.items():
        setattr(faq, field, value)
    faq.save()
    return "updated"


def office_of(row):
    for column in OFFICE_COLUMNS:
        code = OFFICE_NAMES.get(row.get(column.casefold(), "").casefold())
        if code is not None:
            return code
    raise SkippedRow("unknown office")
